import os, re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

md_path = r'C:\Users\Administrator\Desktop\新建文本文档.md'
docx_path = r'C:\Users\Administrator\Desktop\新建文本文档.docx'

doc = Document()

# Set default font
style = doc.styles['Normal']
font = style.font
font.name = '等线'
font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '等线')

with open(md_path, 'r', encoding='utf-8') as f:
    content = f.read()

lines = content.split('\n')
i = 0
in_code_block = False
code_buffer = []

while i < len(lines):
    line = lines[i]
    text = re.sub(r'^\s+\d+\t', '', line).rstrip()
    
    # Code block handling
    if text.strip().startswith('```'):
        if not in_code_block:
            in_code_block = True
            code_buffer = []
        else:
            in_code_block = False
            code_text = '\n'.join(code_buffer)
            if code_text.strip():
                p = doc.add_paragraph()
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x1E, 0x1E, 0x1E)
                # Add shading (grey background)
                from docx.oxml import OxmlElement
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'F2F2F2')
                p.paragraph_format.element.get_or_add_pPr().append(shd)
            code_buffer = []
        i += 1
        continue
    
    if in_code_block:
        code_buffer.append(text)
        i += 1
        continue
    
    # Empty line
    if not text.strip():
        doc.add_paragraph('')
        i += 1
        continue
    
    # Heading: ## Title, ### Title, etc.
    heading_match = re.match(r'^(#{1,6})\s+(.+)$', text)
    if heading_match:
        level = len(heading_match.group(1))
        h_text = heading_match.group(2)
        doc.add_heading(h_text, level=level)
        i += 1
        continue
    
    # Bold: **text**
    # Italic: *text*
    # We handle inline formatting in a simpler way - just add as normal text for now
    doc.add_paragraph(text)
    i += 1

doc.save(docx_path)
print(f'转换完成！文件已保存到: {docx_path}')
print(f'文件大小: {os.path.getsize(docx_path)} bytes')
